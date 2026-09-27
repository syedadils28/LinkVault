import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
import datetime

def export_docx(links, collection_name, filepath):
    """
    Exports a list of links to a professionally formatted Word document.
    """
    document = Document()
    
    # Title
    title = document.add_heading('LINKVAULT', 0)
    title.alignment = 1 # Center
    
    subtitle_text = f"Link Collection: {collection_name}" if collection_name else "Saved Links"
    subtitle = document.add_heading(subtitle_text, 1)
    subtitle.alignment = 1
    
    document.add_paragraph(f"Generated on {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}")
    document.add_page_break()
    
    for i, link in enumerate(links, 1):
        # Link Title
        heading = document.add_heading(f"{i}. {link.title or 'Untitled Link'}", level=2)
        
        # URL
        p_url = document.add_paragraph()
        p_url.add_run("URL: ").bold = True
        p_url.add_run(link.url)
        
        # Type & Collection
        p_meta = document.add_paragraph()
        p_meta.add_run("Type: ").bold = True
        p_meta.add_run(f"{link.link_type}  |  ")
        if link.collection:
            p_meta.add_run("Collection: ").bold = True
            p_meta.add_run(link.collection.name)
            
        # Tags
        if link.tags:
            p_tags = document.add_paragraph()
            p_tags.add_run("Tags: ").bold = True
            p_tags.add_run(", ".join(tag.name for tag in link.tags))
            
        # Description
        if link.description:
            p_desc = document.add_paragraph()
            p_desc.add_run("Description:\n").bold = True
            p_desc.add_run(link.description)
            
        document.add_paragraph("-" * 50)
        
    document.save(filepath)
    return filepath
